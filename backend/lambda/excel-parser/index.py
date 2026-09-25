"""
Excel Parser Lambda
Parses uploaded Excel/Google Sheet workbooks and extracts structured data
for the AI questioning engine. Handles multi-sheet workbooks with financial
models, assumptions, and property data.
"""

import json
import os
import boto3
from io import BytesIO

s3 = boto3.client("s3")
BUCKET = os.environ["UPLOADS_BUCKET"]


def handler(event, context):
    """
    Parse an Excel workbook from S3 and return structured sheet data.

    Expected input:
    {
        "s3Key": "uploads/<sessionId>/workbook.xlsx",
        "sessionId": "...",
        "sheets": ["Summary", "Properties", "Financial Model", "Assumptions"]  // optional filter
    }
    """
    try:
        body = json.loads(event.get("body", "{}"))
        s3_key = body.get("s3Key")
        session_id = body.get("sessionId")
        sheet_filter = body.get("sheets")

        if not s3_key:
            return response(400, {"error": "Missing s3Key"})

        # Download file from S3
        obj = s3.get_object(Bucket=BUCKET, Key=s3_key)
        file_content = obj["Body"].read()

        # Parse Excel workbook
        parsed_data = parse_workbook(file_content, sheet_filter)

        # Store parsed result back to S3 for the AI agent to reference
        if session_id:
            parsed_key = f"sessions/{session_id}/parsed_workbook.json"
            s3.put_object(
                Bucket=BUCKET,
                Key=parsed_key,
                Body=json.dumps(parsed_data, default=str),
                ContentType="application/json",
            )

        return response(
            200,
            {
                "sessionId": session_id,
                "sheets": list(parsed_data.keys()),
                "data": parsed_data,
            },
        )

    except Exception as e:
        print(f"Error parsing workbook: {str(e)}")
        return response(500, {"error": str(e)})


def argb_to_hex(argb_color):
    """Convert openpyxl ARGB color to CSS hex color."""
    if not argb_color or argb_color == "00000000":
        return None
    # ARGB format: first 2 chars are alpha, rest is RGB
    if len(argb_color) == 8:
        return f"#{argb_color[2:]}"  # Skip alpha, return RGB
    elif len(argb_color) == 6:
        return f"#{argb_color}"
    return None


def get_cell_style(cell) -> dict:
    """Extract style information from an openpyxl cell."""
    style = {}

    try:
        # Background color
        fill = cell.fill
        if fill and fill.fgColor and fill.fgColor.type == "rgb" and fill.fgColor.rgb:
            bg_color = argb_to_hex(fill.fgColor.rgb)
            if bg_color and bg_color.lower() != "#ffffff":  # Skip white backgrounds
                style["bg"] = bg_color

        # Font color
        font = cell.font
        if font and font.color and font.color.type == "rgb" and font.color.rgb:
            font_color = argb_to_hex(font.color.rgb)
            if font_color and font_color.lower() != "#000000":  # Skip black font
                style["fc"] = font_color

        # Font style
        if font:
            if font.bold:
                style["b"] = True
            if font.italic:
                style["i"] = True
            if font.underline:
                style["u"] = True

        # Alignment
        align = cell.alignment
        if align:
            if align.horizontal and align.horizontal != "general":
                style["ha"] = align.horizontal  # left, center, right

        # Number format (for detecting percentages, currency, dates)
        if cell.number_format and cell.number_format != "General":
            nf = cell.number_format
            if "%" in nf:
                style["nf"] = "pct"
            elif "$" in nf or "¥" in nf or "€" in nf or "£" in nf:
                style["nf"] = "cur"
            elif "yy" in nf.lower() or "mm" in nf.lower() or "dd" in nf.lower():
                style["nf"] = "date"
    except Exception:
        pass  # Ignore style extraction errors

    return style if style else None


def format_cell_value(value, number_format):
    """
    Apply Excel number format to a value to get the display string.
    This mimics how Excel displays the value.
    """
    if value is None or value == "":
        return ""

    if not isinstance(value, (int, float)):
        return str(value)

    if number_format is None or number_format == "General":
        # For General format, use reasonable precision
        if isinstance(value, float):
            # Check if it's essentially an integer
            if abs(value - round(value)) < 1e-9:
                return str(int(round(value)))
            # For small decimals, show 2 decimal places
            if abs(value) < 1000:
                formatted = f"{value:.2f}"
                # Remove trailing zeros after decimal
                if "." in formatted:
                    formatted = formatted.rstrip("0").rstrip(".")
                return formatted
            else:
                return f"{value:,.0f}"
        return str(int(value)) if isinstance(value, int) else str(value)

    nf = number_format.lower()

    try:
        # Percentage formats
        if "%" in number_format:
            # Value is already in decimal (0.4516 = 45.16%)
            pct_value = value * 100
            # Count decimal places from format
            if "0.00%" in number_format or "0.0%" in number_format:
                decimal_places = number_format.count("0") - 1
                if decimal_places < 0:
                    decimal_places = 0
                return f"{pct_value:.{min(decimal_places, 2)}f}%"
            elif "#,##0%" in number_format or "0%" in number_format:
                return f"{pct_value:.0f}%"
            else:
                return f"{pct_value:.2f}%"

        # Currency formats
        if "$" in number_format:
            if value < 0:
                return f"(${abs(value):,.2f})"
            return f"${value:,.2f}"

        # Number with decimals (0.00, #,##0.00)
        if "0.00" in number_format or "#.00" in number_format:
            return f"{value:,.2f}"

        # Number with one decimal
        if "0.0" in number_format and "0.00" not in number_format:
            return f"{value:,.1f}"

        # Integer/thousands format (#,##0 or 0)
        if "#,##0" in number_format or (nf == "0" and "." not in number_format):
            return f"{value:,.0f}"

        # Default: reasonable decimal display
        if isinstance(value, float):
            if abs(value - round(value)) < 1e-9:
                return str(int(round(value)))
            # Show up to 2 decimal places
            formatted = f"{value:.2f}"
            if "." in formatted:
                formatted = formatted.rstrip("0").rstrip(".")
            return formatted

        return str(value)

    except Exception:
        return str(value)


def parse_workbook(file_content: bytes, sheet_filter: list = None) -> dict:
    """
    Parse Excel workbook bytes into structured JSON.
    Uses openpyxl for .xlsx files.

    Returns values, formulas, AND styles for each sheet.

    Returns dict of sheet_name -> {
        headers: [...],
        rows: [[...], ...],
        formulas: [[...], ...],
        styles: [[{...}, ...], ...],  # Cell styles (bg, fc, b, i, u, ha, nf)
        rowCount: int,
        colWidths: [...],  # Column widths
        mergedCells: [...]  # Merged cell ranges
    }
    """
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter

        # Load twice: once for values, once for formulas/styles
        wb_values = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
        wb_formulas = openpyxl.load_workbook(BytesIO(file_content), data_only=False)
    except ImportError:
        # Fallback: return placeholder indicating openpyxl not available in layer
        return {"error": "openpyxl not available - add to Lambda layer"}

    result = {}
    for sheet_name in wb_values.sheetnames:
        if sheet_filter and sheet_name not in sheet_filter:
            continue

        ws_values = wb_values[sheet_name]
        ws_formulas = wb_formulas[sheet_name]

        # Get all rows including empty ones to preserve structure
        all_rows = []
        all_formulas = []
        all_styles = []
        max_col = ws_values.max_column

        # Use iter_rows on BOTH workbooks simultaneously for performance
        # zip them together so we process each row pair once
        values_rows = list(ws_values.iter_rows(min_row=1, max_row=ws_values.max_row))
        formula_rows = list(
            ws_formulas.iter_rows(min_row=1, max_row=ws_formulas.max_row)
        )

        for val_row, fmt_row in zip(values_rows, formula_rows):
            row_data = []
            formula_row = []
            style_row = []

            for val_cell, fmt_cell in zip(val_row[:max_col], fmt_row[:max_col]):
                cell_value = val_cell.value
                number_format = fmt_cell.number_format

                # Format the value for display
                if cell_value is None:
                    row_data.append("")
                elif isinstance(cell_value, (int, float)):
                    formatted = format_cell_value(cell_value, number_format)
                    row_data.append(formatted)
                else:
                    row_data.append(str(cell_value))

                # Check if cell has a formula
                if (
                    fmt_cell.value
                    and isinstance(fmt_cell.value, str)
                    and fmt_cell.value.startswith("=")
                ):
                    formula_row.append(fmt_cell.value)
                else:
                    formula_row.append("")

                # Extract cell style
                cell_style = get_cell_style(fmt_cell)
                style_row.append(cell_style)

            all_rows.append(row_data)
            all_formulas.append(formula_row)
            all_styles.append(style_row)

        # Extract column widths
        col_widths = []
        for col_idx in range(1, max_col + 1):
            col_letter = get_column_letter(col_idx)
            if col_letter in ws_formulas.column_dimensions:
                width = ws_formulas.column_dimensions[col_letter].width
                if width:
                    col_widths.append(round(width * 7))  # Approximate pixels
                else:
                    col_widths.append(64)  # Default width
            else:
                col_widths.append(64)

        # Extract merged cells
        merged_cells = []
        for merged_range in ws_formulas.merged_cells.ranges:
            merged_cells.append(
                {
                    "start": [merged_range.min_row, merged_range.min_col],
                    "end": [merged_range.max_row, merged_range.max_col],
                }
            )

        # Find the header row (first row with multiple non-empty cells)
        header_idx = 0
        for idx, row in enumerate(all_rows):
            non_empty = sum(1 for cell in row if cell not in [None, "", 0])
            if non_empty >= 3:
                header_idx = idx
                break

        # Generate column headers
        headers = []
        if header_idx < len(all_rows):
            first_row = all_rows[header_idx]
            for col_idx, cell in enumerate(first_row):
                if cell and str(cell).strip():
                    headers.append(str(cell))
                else:
                    col_letter = get_column_letter(col_idx + 1)
                    headers.append(col_letter)
        else:
            headers = [get_column_letter(i + 1) for i in range(max_col)]

        result[sheet_name] = {
            "headers": headers,
            "rows": all_rows,
            "formulas": all_formulas,
            "styles": all_styles,  # NEW: Include styles
            "colWidths": col_widths,  # NEW: Column widths
            "mergedCells": merged_cells,  # NEW: Merged cells
            "rowCount": len(all_rows),
        }

    return result


def response(status_code, body):
    """Format API Gateway response."""
    return {
        "statusCode": status_code,
        "headers": {
            "Content-Type": "application/json",
            "Access-Control-Allow-Origin": "*",
            "Access-Control-Allow-Headers": "Content-Type,Authorization",
        },
        "body": json.dumps(body, default=str),
    }
